from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from src.schemas import Affidavit
from src.mapper import AffidavitMapper


class AffidavitGenerator:

    def __init__(self, affidavit: Affidavit):
        self.affidavit = affidavit
        self.mapped = AffidavitMapper(affidavit).map()

    def generate(self, output_path: str) -> str:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        doc = Document()

        # Default font
        styles = doc.styles
        normal_style = styles["Normal"]
        normal_style.font.name = "Times New Roman"
        normal_style.font.size = Pt(12)

        self._add_centered_bold(doc, self.mapped["court"])
        self._add_centered_bold(doc, self.mapped["jurisdiction"])
        self._add_centered_bold(doc, self.mapped["case_reference"])

        doc.add_paragraph()

        self._add_cause_title(doc)

        doc.add_paragraph()

        self._add_centered_bold(
            doc,
            self.mapped["affidavit_title"]
        )

        doc.add_paragraph()

        # Deponent clause
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run(self.mapped["deponent_clause"])

        doc.add_paragraph()

        # Reply paragraphs
        for paragraph in self.mapped["reply_paragraphs"]:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            p.add_run(
                f"{paragraph['number']}. "
            ).bold = True

            p.add_run(paragraph["content"])

        doc.add_paragraph()

        # Prayer
        prayer_heading = doc.add_paragraph()
        prayer_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        prayer_heading.add_run("PRAYER").bold = True

        for index, prayer in enumerate(self.mapped["prayer"]):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            label = chr(ord("a") + index)

            p.add_run(f"({label}) ").bold = True
            p.add_run(prayer)

        doc.add_paragraph()

        # Attestation
        self._add_attestation(doc)

        doc.add_paragraph()

        # Verification
        verification_heading = doc.add_paragraph()
        verification_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        verification_heading.add_run("VERIFICATION").bold = True

        verification = self.mapped["verification"]

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run(verification["text"])

        p = doc.add_paragraph()
        p.add_run(
            f"Verified at {verification['place']} "
            f"on {verification['date']}."
        )

        doc.add_paragraph()

        # Advocate block
        self._add_advocate_block(doc)

        doc.save(output_path)

        return str(output_path)

    def _add_centered_bold(
        self,
        doc: Document,
        text: str
    ):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(text)
        run.bold = True

    def _add_cause_title(self, doc: Document):
        p = doc.add_paragraph()

        p.add_run(
            f"{self.mapped['petitioner']}"
        ).bold = True

        p.add_run(
            "\n... Petitioner"
        )

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("VERSUS").bold = True

        p = doc.add_paragraph()

        p.add_run(
            f"1. {self.mapped['respondent_1']}"
        )

        p.add_run(
            f"\n2. {self.mapped['respondent_2']}"
        )

        p.add_run(
            "\n... Respondents"
        )

    def _add_attestation(self, doc: Document):
        attestation = self.mapped["attestation"]

        p = doc.add_paragraph()
        p.add_run(
            f"Place: {attestation['place']}"
        )

        p = doc.add_paragraph()
        p.add_run(
            f"Date: {attestation['date']}"
        )

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run("DEPONENT").bold = True

    def _add_advocate_block(self, doc: Document):
        advocate = self.mapped["advocate"]

        p = doc.add_paragraph()

        p.add_run(
            advocate["firm"]
        ).bold = True

        p.add_run(
            f"\nAdvocates for {advocate['representing']}"
        )